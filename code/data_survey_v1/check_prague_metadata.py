"""Inspect TURDATA metadata ZIP without downloading its photos or whole archive."""
from check_prague import Remote
import zipfile,requests,collections
u='https://zenodo.org/api/records/10655033/files/TURDATA_metadata_and_photos.zip/content'
r=requests.head(u,timeout=30);print('HEAD',r.status_code,r.headers.get('Content-Length'))
a=Remote(u);a.n=int(r.headers['Content-Length'])
with zipfile.ZipFile(a) as z:
    items=z.infolist();print('ENTRIES',len(items))
    for x in items:
        if x.filename.lower().endswith(('.xlsx','.txt','.pdf','.csv','.docx')):
            print('DOCUMENT',x.filename,x.file_size,x.compress_size)
            if x.file_size<2_000_000:
                b=z.read(x)
                if x.filename.lower().endswith('.txt'):print('TXT',repr(b[:4000]))
                elif x.filename.lower().endswith('.xlsx'):
                    import openpyxl,io
                    w=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True)
                    for s in w:
                        print('SHEET',s.title,s.max_row,s.max_column)
                        for i,row in enumerate(s.values):
                            vals=[str(v)[:100] for v in row if v is not None]
                            if vals:print('ROW',' | '.join(vals[:12]))
                            if i>100:break
print('FETCHED',a.bytes)
